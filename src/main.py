import argparse
import time

from pywinauto import Desktop
from pywinauto import mouse
from pathlib import Path

from datetime import datetime, date
from openpyxl import load_workbook

TITULO_GERADOR = (
    "Gerador de Relatórios "
    "( Versão: 4.0 ) - Recursos Humanos e Folha de Pagamento"
)

TITULO_ERRO = "Erro Não Previsto"

TITULO_SAIDA_IMPRESSAO = "Salvar saída de impressão como"

TITULO_FICHA = "Ficha_Financeira_Resumo_Geral"

TITULO_IMPRIMIR = "Imprimir"

DIRETORIO_SERVIDORES = Path(
    r"C:\Users\Joao Castro\Desktop\JOAO\Servidores"
)

TEMPO_LIMITE = 60
INTERVALO_VERIFICACAO = 1

def clicar_controle(janela, texto):
    controles = janela.descendants()

    for controle in controles:
        try:
            if controle.window_text() == texto:
                controle.click()
                return True
        except Exception:
            pass

    return False


def encontrar_janela_carregamento():
    desktop = Desktop(backend="win32")

    for janela in desktop.windows():
        try:
            botao_cancelar = janela.child_window(
                title="Cancelar",
                class_name="TButton"
            )

            if botao_cancelar.exists():
                return janela

        except Exception:
            pass

    return None


def encontrar_preview():
    desktop = Desktop(backend="win32")

    for janela in desktop.windows():
        try:
            controle = janela.child_window(
                title="frpPreview",
                class_name="TfrPreview"
            )

            if controle.exists():
                return janela

        except Exception:
            pass

    return None


def tratar_erro():
    desktop = Desktop(backend="win32")

    janela_erro = desktop.window(title=TITULO_ERRO)

    if not janela_erro.exists(timeout=0.2):
        return False

    print("Janela 'Erro Não Previsto' detectada.")

    if not clicar_controle(janela_erro, "&Não"):
        print("Botão 'Não' não encontrado.")
        return False

    print("Clicando em 'Não'.")

    return True


def processar_resultado(nome_funcionario, ano):
    print("Aguardando resultado...")

    janela_carregamento = None
    tela_anterior = None

    while True:

        tela_atual = get_screen_name()

        if tela_atual != tela_anterior:
            print(tela_atual)
            tela_anterior = tela_atual

        # Guarda a janela de carregamento assim que ela aparecer
        if janela_carregamento is None:
            janela_carregamento = encontrar_janela_carregamento()

            if janela_carregamento is not None:
                print("Janela de carregamento encontrada.")

        # Verifica se apareceu erro
        if tratar_erro():
            print("Erro tratado.")

            if janela_carregamento is not None:
                if clicar_controle(janela_carregamento, "Cancelar"):
                    print("Clicando em 'Cancelar'.")
                    return "erro"

                print("Botão 'Cancelar' não encontrado.")
                return None

            print("Janela de carregamento não foi localizada.")
            return None

        # Verifica se abriu o preview
        if tela_atual == "Preview":
            print("Relatório aberto no preview.")

            clicar_botao_preview()

            salvo = imprimir_como_pdf(
                nome_funcionario,
                ano
            )

            if salvo:
                fechar_preview()

            return "preview"

        time.sleep(0.2)


def get_screen_name():
    desktop = Desktop(backend="win32")

    for janela in desktop.windows():
        try:
            titulo = janela.window_text()
            classe = janela.class_name()

            if titulo == TITULO_GERADOR:
                return "Gerador"

            if (
                titulo == TITULO_FICHA
                and classe == "TfrmPreview"
            ):
                return "Preview"

        except Exception:
            pass

    return "Desconhecida"


def alterar_ano(ano, controles):    
    campo_ano = None
    
    for controle in controles:
        try:
            if controle.class_name() != "TcxCustomInnerTextEdit":
                continue

            texto = controle.window_text().strip()

            if texto.isdigit() and 1900 <= int(texto) <= 2100:
                campo_ano = controle
                break

        except Exception:
            pass
        
    if campo_ano is None:
        print("Campo do ano não encontrado.")
        return

    print("Campo do ano encontrado.")
    print("Valor atual:", campo_ano.window_text())
    print("Alterando para:", ano)

    campo_ano.set_focus()

    campo_ano.type_keys("^a")
    campo_ano.type_keys(str(ano))

    print("Valor depois da alteração:", campo_ano.window_text())


def executar_ficha(nome_funcionario, ano, controles):
    botao_ok = None

    for controle in controles:
        try:
            if (
                controle.class_name() == "TBitBtn"
                and controle.window_text() == "&OK"
            ):
                botao_ok = controle
                break
        except Exception:
            pass

    if botao_ok is None:
        print("Botão OK não encontrado.")
        return

    print("Botão OK encontrado.")
    print("Clicando em OK...")

    botao_ok.click()

    print("OK clicado.")

    # Trata o resultado do processamento.
    resultado = processar_resultado(
        nome_funcionario,
        ano
    )

    if resultado == "erro":
        print("Fluxo de erro concluído.")

    elif resultado == "preview":
        fechar_preview()
        print("Processamento concluído.")


def clicar_botao_preview():
    x = 306
    y = 39

    print(f"Clicando no botão em ({x}, {y})...")

    mouse.click(coords=(x, y))

    print("Botão clicado.")


def imprimir_como_pdf(nome_funcionario, ano):
    desktop = Desktop(backend="win32")

    print("Aguardando janela de impressão...")

    janela = None

    while janela is None:
        janelas = desktop.windows(
            class_name="TfrPrintForm"
        )

        for item in janelas:
            try:
                if item.is_visible():
                    janela = item
                    break
            except Exception:
                pass

        if janela is None:
            time.sleep(0.2)

    print("Janela de impressão encontrada.")
    print("Handle:", janela.handle)

    listar_combos_impressao(janela)

    # ---------------------------------------------------------
    # Localiza a ComboBox da impressora
    # ---------------------------------------------------------

    combo_impressora = None

    for controle in janela.descendants():
        try:
            if controle.class_name() != "TComboBox":
                continue

            itens = controle.texts()

            print(
                "Analisando TComboBox:",
                repr(controle.window_text()),
                "| Handle:",
                controle.handle
            )

            if "Microsoft Print to PDF" in itens:
                combo_impressora = controle
                break

        except Exception:
            pass

    if combo_impressora is None:
        print("ComboBox da impressora não encontrada.")
        return False

    # ---------------------------------------------------------
    # Seleciona Microsoft Print to PDF
    # ---------------------------------------------------------

    print(
        "ComboBox da impressora encontrada:",
        combo_impressora.handle
    )

    print(
        "Impressora atual:",
        repr(combo_impressora.window_text())
    )

    print("Selecionando Microsoft Print to PDF...")

    try:
        combo_impressora.select("Microsoft Print to PDF")
    except Exception as e:
        print("Erro ao selecionar impressora:", e)
        return False

    time.sleep(0.05)

    print(
        "Impressora depois da seleção:",
        repr(combo_impressora.window_text())
    )

    # ---------------------------------------------------------
    # Localiza o botão OK
    # ---------------------------------------------------------

    botao_ok = None

    for controle in janela.descendants():
        try:
            if (
                controle.class_name() == "TButton"
                and controle.window_text() == "OK"
            ):
                botao_ok = controle
                break

        except Exception:
            pass

    if botao_ok is None:
        print("Botão OK da impressão não encontrado.")
        return False

    print("Botão OK da impressão encontrado.")
    print("Handle:", botao_ok.handle)
    print("Clicando em OK...")

    botao_ok.click()

    print("OK da impressão clicado.")

    # verifica se a impressão está aberta
    janela = aguardar_janela(TITULO_SAIDA_IMPRESSAO)
    janela.set_focus()

    # Aguarda a janela "Salvar saída de impressão como"
    return salvar_pdf(nome_funcionario, ano)


def salvar_pdf(nome_funcionario, ano):

    print("Aguardando janela para salvar PDF...")

    janela = aguardar_janela(TITULO_SAIDA_IMPRESSAO)
    janela.set_focus()

    print("Janela de salvamento encontrada.")
    print("Título:", janela.window_text())

    pasta_servidor = (
        DIRETORIO_SERVIDORES / nome_funcionario
    )

    pasta_servidor.mkdir(
        parents=True,
        exist_ok=True
    )

    print("Pasta do servidor:", pasta_servidor)

    nome_arquivo = (
        f"FICHA FINANCEIRA {ano}.pdf"
    )

    caminho_arquivo = pasta_servidor / nome_arquivo

    print("Arquivo:", nome_arquivo)
    print("Caminho:", caminho_arquivo)

    # Campo de nome do arquivo
    campo_nome = None

    for controle in janela.descendants():
        try:
            if controle.class_name() != "Edit":
                continue

            rect = controle.rectangle()

            # Campo grande de nome do arquivo
            if rect.width() > 500:
                campo_nome = controle
                break

        except Exception:
            pass

    if campo_nome is None:
        print("Campo do nome do arquivo não encontrado.")
        return False

    # Define o caminho completo
    campo_nome.set_focus()
    campo_nome.set_edit_text(str(caminho_arquivo))

    print("Nome do arquivo preenchido.")

    # Botão Salvar
    botao_salvar = None

    for controle in janela.descendants():
        try:
            if (
                controle.class_name() == "Button"
                and controle.window_text() == "Sa&lvar"
            ):
                botao_salvar = controle
                break
        except Exception:
            pass

    if botao_salvar is None:
        print("Botão Salvar não encontrado.")
        return False

    print("Clicando em Salvar...")
    botao_salvar.click()

    print("Comando de salvar enviado.")

    confirmar_sobrescrita_pdf()

    print("Processo de salvamento concluído.")
    return True


def listar_combos_impressao(janela):
    print("\nCOMBOBOXES ENCONTRADAS")
    print("=" * 70)

    combos = janela.descendants(class_name="TComboBox")

    print(f"Quantidade: {len(combos)}\n")

    for i, combo in enumerate(combos):
        try:
            print("-" * 70)
            print(f"COMBO [{i}]")
            print("Handle :", combo.handle)
            print("Texto  :", repr(combo.window_text()))
            print("Classe :", combo.class_name())

            try:
                itens = combo.texts()

                print("Quantidade de itens:", len(itens))

                for j, item in enumerate(itens):
                    print(f"  [{j}] {item!r}")

            except Exception as e:
                print("Não foi possível obter os itens:", e)

        except Exception as e:
            print(f"Erro na ComboBox [{i}]: {e}")

def selecionar_impressora(janela, nome_impressora):
    combos = janela.descendants(
        class_name="TComboBox"
    )

    for combo in combos:
        try:
            itens = combo.texts()

            if nome_impressora in itens:
                print("ComboBox da impressora encontrada.")
                print("Handle:", combo.handle)
                print("Impressora atual:", repr(combo.window_text()))

                print("Selecionando:", nome_impressora)

                combo.select(nome_impressora)

                time.sleep(0.5)

                print(
                    "Impressora depois da seleção:",
                    repr(combo.window_text())
                )

                return True

        except Exception as e:
            print(
                f"Erro ao analisar ComboBox "
                f"{combo.handle}: {e}"
            )

    print(
        f"Impressora {nome_impressora!r} "
        "não encontrada."
    )

    return False


def obter_nome_funcionario(janela):
    listboxes = janela.descendants(
        class_name="TListBox"
    )

    for listbox in listboxes:
        try:
            for item in listbox.texts():
                item = item.strip()

                if " - " not in item:
                    continue

                codigo, nome = item.split(" - ", 1)

                if codigo.strip().isdigit() and nome.strip():
                    return nome.strip()

        except Exception:
            pass
        
    print("Nome do funcionário não encontrado.")
    return input("Digite o nome a ser considerado: ")


def fechar_preview():

    janela = aguardar_janela(TITULO_FICHA)
    janela.set_focus()

    print("Fechando Ficha_Financeira_Resumo_Geral...")

    janela.close()

    print("Preview fechado.")
    return True


def preencher_matricula(matricula, controles):

    janela = aguardar_janela(TITULO_GERADOR)
    janela.set_focus()

    print("Janela encontrada para preenchimento da matrícula.")

    # ---------------------------------------------------------
    # Localiza a aba "Funcionário"
    # ---------------------------------------------------------

    aba_funcionario = None

    for controle in controles:
        try:
            if (
                controle.class_name() == "TcxTabSheet"
                and controle.window_text().strip() == "Funcionário"
            ):
                aba_funcionario = controle
                break
        except Exception:
            pass

    if aba_funcionario is None:
        print("Aba 'Funcionário' não encontrada.")
        return False

    print(
        "Aba 'Funcionário' encontrada:",
        aba_funcionario.handle
    )

    # ---------------------------------------------------------
    # Localiza os TelLookupCombo dentro da aba
    # ---------------------------------------------------------

    combos = aba_funcionario.descendants(
        class_name="TelLookupCombo"
    )

    print(
        f"TelLookupCombo encontrados na aba: {len(combos)}"
    )

    for combo in combos:
        try:
            rect = combo.rectangle()

            print(
                f"TelLookupCombo | "
                f"texto={combo.window_text()!r} | "
                f"posição=({rect.left},{rect.top}) | "
                f"tamanho=({rect.width()}x{rect.height()})"
            )

        except Exception:
            pass

    # ---------------------------------------------------------
    # O campo Funcionário é o TelLookupCombo que está
    # aproximadamente em X=465 / Y=554.
    #
    # Os outros TelLookupCombo ficam nos campos de
    # estabelecimento/lotação acima.
    #
    # Não usamos handle.
    # ---------------------------------------------------------

    if not combos:
        print("Nenhum TelLookupCombo encontrado na aba 'Funcionário'.")
        return False

    if len(combos) > 1:
        print(
            "ATENÇÃO: foram encontrados vários TelLookupCombo "
            "na aba 'Funcionário'."
        )

    campo_matricula = combos[0]

    print()
    print("Campo de matrícula encontrado.")
    print("Classe:", campo_matricula.class_name())
    print("Posição:", campo_matricula.rectangle())

    print()
    print("Campo de matrícula encontrado.")
    print("Classe:", campo_matricula.class_name())
    print("Posição:", campo_matricula.rectangle())

    # ---------------------------------------------------------
    # Preenche matrícula
    # ---------------------------------------------------------

    campo_matricula.set_focus()

    time.sleep(0.1)

    campo_matricula.type_keys(
        "^a"
    )

    campo_matricula.type_keys(
        str(matricula)
    )

    print(
        "Matrícula preenchida:",
        matricula
    )

    # ---------------------------------------------------------
    # CONFIRMA COM ENTER
    # ---------------------------------------------------------

    campo_matricula.type_keys(
        "{ENTER}"
    )

    print("ENTER enviado para a matrícula.")

    return True


def confirmar_sobrescrita_pdf():

    print()
    print("Verificando confirmação de sobrescrita...")

    desktop = Desktop(backend="win32")

    limite = time.time() + 5

    while time.time() < limite:

        try:

            janela = desktop.window(
                class_name="#32770",
                title="Confirmar Salvar como"
            )

            if janela.exists():

                print(
                    'Janela "Confirmar Salvar como" encontrada.'
                )

                # Procura o botão &Sim
                botoes = janela.descendants(
                    class_name="Button"
                )

                for botao in botoes:

                    try:

                        texto = botao.window_text().strip()

                        if texto == "&Sim":

                            print(
                                'Botão "&Sim" encontrado.'
                            )

                            botao.click()

                            print(
                                'Botão "&Sim" clicado.'
                            )

                            return True

                    except Exception:
                        pass

                print(
                    'Janela encontrada, mas botão "&Sim" '
                    'não foi localizado.'
                )

        except Exception:
            pass

        time.sleep(0.1)

    print(
        'Janela "Confirmar Salvar como" não apareceu.'
    )

    return False



def ler_vinculos(caminho_excel):
    caminho_excel = Path(caminho_excel)

    if not caminho_excel.exists():
        raise FileNotFoundError(
            f"Arquivo de vínculos não encontrado: {caminho_excel}"
        )

    wb = load_workbook(
        caminho_excel,
        read_only=True,
        data_only=True
    )

    try:
        ws = wb.active
        vinculos = []

        for numero_linha, linha in enumerate(
            ws.iter_rows(min_col=1, max_col=3, values_only=True),
            start=1
        ):
            matricula, inicio, fim = linha

            # Ignora linhas vazias
            if matricula is None and inicio is None and fim is None:
                continue

            # Converte a matrícula
            try:
                matricula_num = int(matricula)
            except (TypeError, ValueError, OverflowError):
                # Ignora cabeçalhos ou linhas inválidas
                continue

            def converter_data(valor):
                if isinstance(valor, datetime):
                    return valor.date()

                if isinstance(valor, date):
                    return valor

                if isinstance(valor, str):
                    return datetime.strptime(
                        valor.strip(), "%d/%m/%Y"
                    ).date()

                raise ValueError(
                    f"Data inválida na linha {numero_linha}: {valor!r}"
                )

            data_inicio = converter_data(inicio)
            data_fim = converter_data(fim)

            if data_inicio > data_fim:
                raise ValueError(
                    f"Data inicial posterior à final na linha "
                    f"{numero_linha}."
                )

            vinculos.append({
                "matricula": matricula_num,
                "inicio": data_inicio,
                "fim": data_fim,
            })

        if not vinculos:
            raise ValueError(
                "Nenhum vínculo válido encontrado nas colunas A, B e C."
            )

        return vinculos

    finally:
        wb.close()


def localizar_janela(titulo_esperado):
    """
    Localiza uma janela pelo título ou por um trecho dele.
    Retorna um objeto pywinauto ou None.
    """
    for janela in Desktop(backend="win32").windows():
        try:
            titulo = janela.window_text().strip()

            if titulo_esperado.casefold() in titulo.casefold():
                print(f"[OK] Janela identificada: {titulo}")
                return janela

        except Exception:
            continue

    return None


def aguardar_janela(titulo_esperado, timeout=TEMPO_LIMITE):
    """
    Aguarda até a janela aparecer.
    Retorna o objeto da janela ou None se atingir o limite.
    """
    print(f"[INFO] Aguardando a janela: {titulo_esperado}")

    limite = time.monotonic() + timeout

    while time.monotonic() < limite:
        janela = localizar_janela(titulo_esperado)

        if janela is not None:
            return janela

        time.sleep(INTERVALO_VERIFICACAO)

    print(
        f"[ERRO] A janela '{titulo_esperado}' "
        f"não apareceu em {timeout} segundos."
    )
    return None


def main():
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "entrada",
        nargs="?",
        help="Caminho do vinculos.xlsx ou intervalo de anos"
    )

    parser.add_argument(
        "--mat",
        type=int,
        help="Matrícula para processamento manual"
    )

    parser.add_argument(
        "--padrao",
        action="store_true",
        help="Executa matrícula 3198 no ano 2026"
    )

    args = parser.parse_args()

    print("Esperando janela: {TITULO_GERADOR}\n")
    janela = aguardar_janela(TITULO_GERADOR)
    janela.set_focus()

    print("Janela encontrada.")

    controles = janela.descendants()
    
    # ========================================================
    # MODO PADRÃO
    # ========================================================

    if args.padrao:
        matricula = 3198
        ano_inicial = 2026
        ano_final = 2026

        print("MODO PADRÃO")
        print(f"Matrícula: {matricula}")
        print(f"Ano: {ano_inicial}")

        if not preencher_matricula(matricula, controles):
            print("Não foi possível preencher a matrícula.")
            return

        alterar_ano(ano_inicial, controles)

        nome_funcionario = obter_nome_funcionario(janela)
        print("Funcionário:", nome_funcionario)
        
        executar_ficha(nome_funcionario, ano_inicial, controles)
        return

    
    # ========================================================
    # MODO AUTOMÁTICO: ARQUIVO VINCULOS.XLSX
    # ========================================================

    if args.entrada and args.entrada.lower().endswith(".xlsx"):
        vinculos = ler_vinculos(args.entrada)

        print("=" * 70)
        print("ETAPA 1: PREENCHER TODAS AS MATRÍCULAS")
        print("=" * 70)

        # Primeiro, preencher todas as matrículas
        for indice, vinculo in enumerate(vinculos, start=1):
            matricula = vinculo["matricula"]

            print(
                f"[{indice}/{len(vinculos)}] "
                f"Preenchendo matrícula: {matricula}"
            )

            if not preencher_matricula(matricula, controles):
                print(
                    f"ERRO ao preencher a matrícula {matricula}. "
                    "Processamento interrompido."
                )
                return

        print()
        print("=" * 70)
        print("ETAPA 2: PROCESSAR OS ANOS DE CADA VÍNCULO")
        print("=" * 70)

        # Depois, percorrer todos os vínculos e seus anos
        for indice, vinculo in enumerate(vinculos, start=1):
            matricula = vinculo["matricula"]
            ano_inicial = vinculo["inicio"].year
            ano_final = vinculo["fim"].year

            print()
            print("#" * 70)
            print(f"VÍNCULO {indice}/{len(vinculos)}")
            print(f"Matrícula: {matricula}")
            print(f"Anos: {ano_inicial}-{ano_final}")
            print("#" * 70)

            for ano in range(ano_inicial, ano_final + 1):
                print(f"Processando matrícula {matricula}, ano {ano}...")
                alterar_ano(ano, controles)

                nome_funcionario = obter_nome_funcionario(janela)
                print("Funcionário:", nome_funcionario)

                
                executar_ficha(nome_funcionario, ano, controles)

        print()
        print("=" * 70)
        print("PROCESSAMENTO FINALIZADO.")
        print("=" * 70)
        return


    # ========================================================
    # MODO MANUAL: MATRÍCULA + INTERVALO DE ANOS
    # Exemplo: python src/main.py --mat 22057 2017-2018
    # ========================================================

    if args.mat is None:
        parser.error(
            "Informe o caminho de um vinculos.xlsx, "
            "--mat com o intervalo, ou --padrao."
        )

    if not args.entrada:
        parser.error(
            "Informe o intervalo de anos. Exemplo: 2017-2024"
        )

    try:
        ano_inicial, ano_final = map(
            int, args.entrada.split("-")
        )
    except ValueError:
        parser.error(
            "Intervalo inválido. Use o formato ANO_INICIAL-ANO_FINAL."
        )

    if ano_inicial > ano_final:
        parser.error(
            "O ano inicial não pode ser maior que o ano final."
        )

    if not preencher_matricula(args.mat):
        print("Não foi possível preencher a matrícula.")
        return

    for ano in range(ano_inicial, ano_final + 1):
        print()
        print("=" * 70)
        print(f"MATRÍCULA: {args.mat}")
        print(f"PROCESSANDO ANO: {ano}")
        print("=" * 70)

        alterar_ano(ano, controles)

        nome_funcionario = obter_nome_funcionario(janela)
        print("Funcionário:", nome_funcionario)
        
        executar_ficha(nome_funcionario, ano, controles)


if __name__ == "__main__":
    main()